from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import json, re, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parent
SRC=Path(r'C:\Users\kirby\Downloads\LMS FINAL (1).docx')
OUT=Path(r'C:\Users\kirby\Desktop\lms-monorepo-skeleton\output\documents')
OUT.mkdir(parents=True,exist_ok=True)
FINAL=OUT/'CdM_LMS_Manuscript_Revised_Comparison.docx'
source=Document(SRC);d=Document(SRC)
C=json.loads((ROOT/'content.json').read_text(encoding='utf-8'))
REFS=json.loads((ROOT/'references.json').read_text(encoding='utf-8'))
source_paras=[p.text.strip() for p in source.paragraphs]
cover=[deepcopy(p._p) for p in source.paragraphs[:26]]
for child in list(d._element.body):
    if child.tag!=qn('w:sectPr'):d._element.body.remove(child)
for el in cover:d._element.body.insert(len(d._element.body)-1,el)
cover_title=d.paragraphs[3]
for r in list(cover_title.runs):r._r.getparent().remove(r._r)
r=cover_title.add_run(C['title']);r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=True
cover_title.style=d.styles['Title'];cover_title.alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_comment(runs=[r],text='Confirm the formally approved title. The earlier project instruction specified the multi-platform title used consistently in this revision.',author='Manuscript Review',initials='MR')
datep=next(p for p in d.paragraphs if p.text.strip()=='May 2026')
d.add_comment(runs=datep.runs,text='The supplied cover date has been preserved. Confirm the actual submission date before filing. Later-2026 sources whose dates would conflict with this cover were excluded from the revised literature.',author='Manuscript Review',initials='MR')

def clean(text):
    text=text.replace('BSCPE','BSCpE').replace('Pan et al. (2023)','Pan et al. (2022)').replace('Jack et al. (2025)','Jack (2025)')
    verbs={'describes':'describe','emphasizes':'emphasize','identifies':'identify','provides':'provide','supports':'support'}
    text=re.sub(r'(et al\. \(\d{4}\)) (describes|emphasizes|identifies|provides|supports)\b',lambda m:m[1]+' '+verbs[m[2]],text)
    text=text.replace('points and leaderboards in a flipped statistics classroom and discussed benefits','points and leaderboards in a flipped statistics classroom and considered benefits')
    return text
def para(text='',bold=False,indent=True,style='Normal',before=12,after=12,spacing=2.0):
    p=d.add_paragraph(style=style);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    pf=p.paragraph_format;pf.first_line_indent=Inches(.5) if indent else Inches(0)
    pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=spacing;pf.widow_control=True
    r=p.add_run(clean(text));r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0);r.bold=bold
    r._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Times New Roman')
    return p
def heading(text,level=1,page=False):
    p=para(text,True,False,'Title' if level==0 else f'Heading {min(level,3)}',before=18,after=10,spacing=1.0)
    p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True;p.paragraph_format.page_break_before=page
    if level==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p
def comment(p,text):d.add_comment(runs=p.runs,text=text,author='Manuscript Review',initials='MR')
def section(start,end):
    i=next(i for i,t in enumerate(source_paras) if t.startswith(start));j=next((j for j in range(i+1,len(source_paras)) if source_paras[j].startswith(end)),len(source_paras))
    return [t for t in source_paras[i+1:j] if t]
def paragraphs(items):
    return [para(t) for t in items]
fig_num=0
def figure(name,title,width=6.4,height=8.4,page=False):
    global fig_num
    fig_num+=1
    cap=para(f'Figure {fig_num}. {title}',False,False,before=0,after=8,spacing=1.0)
    cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.keep_with_next=True;cap.paragraph_format.page_break_before=page
    for r in cap.runs:r.italic=True;r.font.size=Pt(11)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1
    path=ROOT/'figures'/f'{name}.png';im=Image.open(path);ratio=im.height/im.width
    w=min(width,height/ratio)
    pic=p.add_run().add_picture(str(path),width=Inches(w))
    pic._inline.docPr.set('descr',title)
    return cap
table_num=0
def table(title,headers,rows,widths):
    global table_num
    table_num+=1
    cap=para(f'Table {table_num}. {title}',False,False,before=10,after=8,spacing=1.0)
    cap.paragraph_format.keep_with_next=True
    for r in cap.runs:r.italic=True;r.font.size=Pt(11)
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for vals in [headers]+rows:
        cells=t.rows[0].cells if vals is headers else t.add_row().cells
        for i,(cell,text) in enumerate(zip(cells,vals)):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');marg.append(e)
            tcpr.append(marg)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'B8B8B8');borders.append(e)
            tcpr.append(borders)
            if vals is headers:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');tcpr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_with_next=False
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(clean(str(text)));r.font.name='Times New Roman';r.font.size=Pt(10.5);r.bold=vals is headers;r.font.color.rgb=RGBColor(0,0,0)
        trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if vals is headers:trpr.append(OxmlElement('w:tblHeader'))
    # Keep a table header and its first row together, without keeping the entire table.
    for p in t.rows[0].cells[0].paragraphs:p.paragraph_format.keep_with_next=True
    return cap

# Repair the source header's floating frame while preserving its logos and fonts.
sec=d.sections[0];header=sec.header
for el in list(header._element):header._element.remove(el)
hp=header.add_paragraph();hp.paragraph_format.space_after=Pt(0);hp.paragraph_format.space_before=Pt(0);hp.paragraph_format.line_spacing=Pt(1)
ht=header.add_table(rows=1,cols=3,width=Inches(6.6875));ht.alignment=WD_TABLE_ALIGNMENT.CENTER;ht.autofit=False
for i,w in enumerate([.85,4.9875,.85]):ht.columns[i].width=Inches(w);ht.cell(0,i).width=Inches(w)
for cell in ht.rows[0].cells:
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
ht.cell(0,0).paragraphs[0].add_run().add_picture(str(ROOT/'media'/'image2.png'),width=Inches(.78))
ht.cell(0,2).paragraphs[0].add_run().add_picture(str(ROOT/'media'/'image1.jpg'),width=Inches(.78))
center=ht.cell(0,1).paragraphs[0]
for text,family,size,bold in [('Colegio de Montalban','Cambria',22,True),('Kasiglahan Village, San Jose, Rodriguez, Rizal','Courier New',8,False),('Institute of Computing Studies','Cambria',19,True)]:
    if center.runs:center.add_run().add_break()
    rr=center.add_run(text);rr.font.name=family;rr.font.size=Pt(size);rr.bold=bold;rr.font.color.rgb=RGBColor(0,0,0)
hp=header.add_paragraph();hp.alignment=WD_ALIGN_PARAGRAPH.CENTER;hp.paragraph_format.space_after=Pt(0);hp.paragraph_format.space_before=Pt(0);hp.paragraph_format.line_spacing=1
hp.add_run().add_picture(str(ROOT/'media'/'image11.png'),width=Inches(6.6875),height=Inches(.04))

heading('CHAPTER 1 INTRODUCTION',0,True)
heading('1.1 BACKGROUND OF THE STUDY')
background=section('1.1 BACKGROUND','1.2 STATEMENT')
background[0]='LMS research connects course access with planning, monitoring, and reflection. Du et al. (2025) examined self-regulated learning support and reported stronger self-regulation, with an examination advantage on one test but not another. Çakiroğlu et al. (2024) related LMS interaction profiles to self-regulation. These findings support examining connected learning tools without assuming that access alone improves achievement.'
background[1]='Computing education also requires engagement mechanisms with clear instructional purposes. Ahmed et al. (2025) evaluated a gamified mobile/cloud LMS in an Object-Oriented Programming course and reported favorable achievement and engagement findings in that setting. Jaramillo-Mediavilla et al. (2024) reviewed gamification evidence and identified the need to examine motivation and performance together. Those findings justify investigating participation support but cannot be presumed for CdM LMS.'
paragraphs(background)
heading('1.2 STATEMENT OF THE PROBLEM')
for text in section('1.2 STATEMENT','1.3 Statement'):
    text=text.replace('What is the level of the LMS in supporting','What is the evaluated software quality of CdM LMS in supporting').replace('What is the level of the system in','What is the evaluated software quality of CdM LMS in')
    para(text,indent=not bool(re.match(r'^\d\.|^•',text)))
heading('1.3 OBJECTIVES OF THE STUDY')
heading('1.3.1 General Objective',2);para(section('1.3.1 General','1.3.2 Specific')[0])
heading('1.3.2 Specific Objectives',2)
for text in section('1.3.2 Specific','1.4 CONCEPTUAL'):
    if text.startswith('The study'):continue
    if re.match(r'^\d+\.',text):
        text=text.replace('verified and approved accounts','account-status controls').replace('Develop the proposed faculty directory','Develop the faculty directory')
        para(text,indent=False)
heading('1.4 CONCEPTUAL FRAMEWORK',page=True);figure('conceptual','Conceptual framework of CdM LMS using IPO with feedback',height=8.0)
paragraphs(C['conceptual'])
heading('1.5 OPERATIONAL FRAMEWORK');figure('agile','Adapted Agile development cycle for CdM LMS',height=4.2)
paragraphs(C['operational'])
heading('1.6 SCOPE AND LIMITATIONS');paragraphs(C['scope'])
heading('1.7 SIGNIFICANCE OF THE STUDY')
for text in section('1.7 SIGNIFICANCE','1.8 GENERAL'):
    text=text.replace('The developed Colegio de Montalban Learning Management System (CdM LMS) is significant','The proposed CdM LMS may contribute').replace('as it provides','by providing')
    if text.startswith('Academic supervisors'):
        text='Super Users and Program Heads. Restricted Super User supervision may support authorized academic oversight. Program Heads may benefit from faculty information and designation-specific event publication. These responsibilities will remain distinct from teaching and system administration.'
    para(text)
heading('1.8 GENERAL FEATURES OF THE SYSTEM')
p=para('CdM LMS comprises complementary applications connected through shared backend services and a centralized database. The following groups describe intended feature coverage. A release record will distinguish existing modules, planned extensions, tested functions, and functions enabled for evaluation. Availability is governed by role and platform coverage.')
comment(p,'Before evaluation, replace the proposed release-status record with confirmed module/build evidence. The academic coordination extension remains planned in the inspected repository; its inclusion here is a development requirement, not a completed result.')
for title,text in C['features']:heading(title,2);para(text)
table('Role and platform responsibilities',['Role','Principal interface','Permitted responsibility','Boundary'],[
 ['Student','Web; selected mobile functions','Approved coursework, own grades, study tools and collaboration','No faculty directory or supervision'],
 ['Instructor','Web; selected mobile functions','Assigned teaching, grading, enrollment and schedules','Program Head capabilities require designation'],
 ['Super User','Web academic operations hub','Safe supervision and narrow faculty coordination','No grading, assessments, schedule editing or system administration'],
 ['Administrator','Native Desktop Console','Protected account/configuration, audit and supported continuity operations','Additional browser/mobile administration is outside this evaluation scope']
],[1.0,1.55,2.5,1.63])
heading('1.9 DEFINITION OF TERMS')
definitions=section('1.9 DEFINITION','CHAPTER 2')
definitions=[t for t in definitions if not t.startswith('The following')]
definitions=[t.replace('Spaced-Repetition-Oriented Review - Flashcard review organized to support repeated study over time, without assuming an unverified scheduling algorithm.','Flashcard Review - Repeated study of saved question-and-answer cards using the review and mastery functions available in the evaluated release.') for t in definitions]
definitions.extend(['BSIT - Bachelor of Science in Information Technology, one of the ICS programs within the study context.','BSCpE - Bachelor of Science in Computer Engineering, one of the ICS programs within the study context.','Dean - An academic leadership designation represented in the institute directory; the designation itself does not grant additional system permissions.'])
for text in sorted(definitions,key=lambda t:t.casefold()):para(text,indent=False,before=6,after=6)

heading('CHAPTER 2 REVIEW OF RELATED LITERATURE',0,True)
intro=section('CHAPTER 2','2.1.1 Multi')
for text in intro:
    text=text.replace('seven topics','eight topic groups')
    if text.startswith('The project developers reviewed'):
        text=text.replace('and Online Collaboration and Knowledge Sharing.','Online Collaboration and Knowledge Sharing; and Academic Coordination and Information Boundaries.')
    para(text)
heading('2.1 Related Literature')
start=source_paras.index(next(t for t in source_paras if t.startswith('2.1.1 Multi')))
end=source_paras.index(next(t for t in source_paras if t.startswith('CHAPTER 3')))
inserted=False
for text in source_paras[start:end]:
    if not text:continue
    if text.startswith('2.2 ') and not inserted:
        heading('2.1.8 Academic Coordination and Information Boundaries',2)
        para('Institutional LMS support extends beyond individual coursework. Guico (2026) identifies faculty preparation, resource access, and continuing training as management concerns. Llamas et al. (2025) identifies monitoring, assessment, communication, and course organization as user-experience concerns. Together, these findings provide reasons to examine faculty coordination and access to appropriate course information. They do not validate the particular schedule schema or supervision policy proposed for CdM LMS.')
        para('The project’s academic coordination requirements will therefore be established through local consultation and explicit role policies. Directory information, timetable revisions, event audiences, and restricted supervision will be tested against those requirements. Academic designation and technical privilege will remain separate; shared calendar visibility will not imply access to private notes or individual assessment records.')
        inserted=True
    if text.startswith('2.5 Synthesis'):text=text.replace('2.5','2.4',1)
    if text.startswith('2.6 Research'):text=text.replace('2.6','2.5',1)
    if re.match(r'^2\.\d',text):heading(text,2 if text.count('.')>=2 else 1);continue
    if 'Bernardez et al.' in text:continue
    if 'Francisco (2026)' in text:
        if text.startswith('Institutional support'):
            text=text.split('Francisco (2026)')[0].strip()+' These findings place course management within faculty support and institutional coordination.'
        elif text.startswith('The reviewed Philippine'):text='The reviewed Philippine studies inform the project’s computing-course, LMS, programming, and collaborative workflows. Their findings do not establish conditions within ICS. Local requirements gathering will identify the needs and experiences of the intended users.'
        else:text=text.replace(', and Francisco (2026)','').replace(' and Francisco (2026)','')
    text=text.replace('these five studies','the reviewed studies').replace('These five studies','The reviewed studies').replace('The five studies','The reviewed studies').replace('Across these six studies','Across the reviewed studies').replace('Together, the five studies','Together, the reviewed studies').replace('The standard is a supporting framework rather than one of the five studies counted within each topic.','The standard supports the quality criteria used in this study.')
    if 'Jack et al. (2025)' in text:text=text.replace('Jack et al. (2025) examined','Jack (2025), in a conference abstract, discussed')
    if text.startswith('The reviewed studies address digital learning'):
        chunks=text.split('A second part');para(chunks[0].strip())
        if len(chunks)>1:
            rest='A second part'+chunks[1];chunks=rest.split('A third part');para(chunks[0].strip())
            if len(chunks)>1:para('A third part'+chunks[1])
    else:para(text)

heading('CHAPTER 3 METHODOLOGY',0,True)
for no,title,key in [('3.1','Research Design','design'),('3.2','System and Project Overview','overview'),('3.3','Local Requirements Gathering','requirements'),('3.4','Respondents','respondents'),('3.5','Sampling Technique','sampling'),('3.6','Instruments and Data Collection','instruments')]:
    heading(no+' '+title);ps=paragraphs(C[key])
    if key=='requirements':comment(ps[0],'Confirm the planned consultation method, permissions, and recruitment arrangements with the adviser. This is a proposed procedure, not a report that interviews have already occurred.')
    if key=='respondents':comment(ps[-1],'Supply justified target numbers for students, instructors and IT experts, invitations, and pilot participants before recruitment. No sample size was invented. Confirm the proposed number and eligibility of validators.')
table('Role appropriate evaluation tasks',['Evaluator','Account/platform','Representative tasks'],[
 ['ICS students','Student; web/selected mobile','Approved coursework, own grades, notes/tasks, study tools and Knowledge Exchange'],
 ['ICS instructors','Instructor; web/selected mobile','Assigned materials, submissions, grading, enrollment and schedules'],
 ['Eligible Program Heads','Instructor designation; web','Permitted faculty/program event publication'],
 ['Qualified IT experts','Controlled Super User and Administrator accounts','Restricted supervision, protected console operations and direct technical checks']
],[1.35,2.0,3.33])
heading('3.7 Technical Testing and Acceptance');ps=paragraphs(C['testing']);comment(ps[1],'Before testing, approve and record dataset size, active sessions, request rate, simultaneous executions, repetitions, response-time thresholds, and resource-observation tools. These cannot be inferred from registered users or hardware estimates.')
table('Technical evidence for selected quality characteristics',['Characteristic','Checks','Evidence'],[
 ['Functional Suitability','Controlled computations, submissions, schedules and conversions','Expected/actual output and coverage'],
 ['Interaction Capability','Role-appropriate tasks, navigation and feedback','Completion observations and ratings'],
 ['Security','Prohibited API/action requests, private-record isolation','Access decisions and audit evidence'],
 ['Performance Efficiency','Recorded workloads, timings and resource use','Measurements with setup and repetitions'],
 ['Reliability','Interrupt/retry, duplicate/stale save, persistence','Record consistency and recovery outcomes'],
 ['Compatibility','Permitted shared records; console co-existence','Client consistency and workstation observations']
],[1.5,2.75,2.43])
heading('3.8 Data Analysis and Statistical Treatment');ps=paragraphs(C['analysis']);comment(ps[1],'Adviser/statistician review is required for the proposed scoring and characteristic aggregation. Retain distributions and valid counts. No unapproved interpretation bands or inferential claims have been inserted.')
heading('3.9 Participant Protection and Information Handling');ps=paragraphs(C['ethics']);comment(ps[-1],'Complete the approved retention period, storage location, access list, withdrawal handling, and institutional permission details. Do not treat this planned protocol as proof that permissions have already been granted.')
heading('3.10 Hardware and Deployment Requirements')
para('Equipment specifications below are planning estimates for development and a small evaluation pilot. They are not experimentally established minimums or verified concurrent capacity. The actual device and hosting configuration will be recorded before task sessions.')
table('Estimated development and client hardware',['Equipment','Estimated baseline','Recommended','Purpose'],[
 ['Development workstation','4 cores; 16 GB RAM; 50 GB free SSD','6+ cores; 32 GB RAM; 100 GB free SSD','Build clients and run local checks'],
 ['Web computer','Dual core; 4 GB RAM; 1366×768; 1 GB free','4 cores; 8 GB RAM; 1080p; 5 GB free','Academic and coordination tasks'],
 ['Android device','Build-compatible; about 4 GB RAM; 1 GB free','6+ GB RAM; 2 GB free','Enabled mobile tasks'],
 ['iOS device if evaluated','Compatible iPhone/iPad; 1 GB free','Supported device; 2 GB free','Verified iOS build only'],
 ['Administrator workstation','64-bit dual core; 4 GB RAM; 2 GB free','4 cores; 8 GB RAM; 5 GB free','Native Console and exports'],
 ['Network equipment','Reliable router/AP and connectivity','Wired laboratory links and sufficient Wi-Fi capacity','Connect authorized clients/services']
],[1.45,1.8,1.8,1.63])
para('Storage allowances cover dependencies, temporary files, cache, and downloads rather than measured installer sizes. Emulators, containers, and datasets may need more development space. Enabled 3D activities require a tested WebGL 2-capable browser/graphics configuration; dedicated graphics is not assumed for ordinary LMS tasks.')
table('Proposed pilot hosting and network configuration',['Component','Initial estimate','Condition to verify'],[
 ['Application service','2 vCPU / 4 GB RAM on a conventional server','For serverless hosting, record plan/function limits instead'],
 ['PostgreSQL','2 vCPU / 4 GB RAM; 10–20 GB initial records/index storage','Connection limits, growth and query behavior'],
 ['Persistent file storage','20–50 GB initial object-storage allowance','Materials, submissions, snapshots and retention'],
 ['Code execution if self-hosted','Separate 4 vCPU / 8 GB RAM; 30–50 GB storage','Languages, isolation, test limits and execution concurrency'],
 ['Backups','Separate protected storage','Retention capacity and demonstrated restoration'],
 ['Individual connection','About 5 Mbps down / 2 Mbps up; prefer 10/5','Measured latency, contention and file sizes'],
 ['Shared lab connection','About 100 Mbps down / 20–50 Mbps up for an initial 20–30-device session','Actual simultaneous workload and connection stability']
],[1.45,2.55,2.68])
para('The initial sizing assumption is approximately 25–50 simultaneously active ordinary-task users, not a research sample size or proven capacity. Code executions will be limited and measured separately. Lower bandwidth may support basic retrieval, while uploads and concurrent activity may need more. Managed-service quotas and the demonstrated workload take precedence over these estimates.')
heading('3.11 Software and Client Requirements')
para('The software record identifies the current development stack and proposed evaluation targets. It will be frozen against the actual release revision; framework support alone does not establish CdM client compatibility.')
table('Software environment and release configuration',['Software/service','Recorded choice','Purpose/qualification'],[
 ['Development OS','Supported Windows/macOS/Linux environment','Platform-specific build requirements apply'],
 ['Node and workspace tools','Supported Node 22 LTS patch ≥22.12; pnpm 10.32.1; Turborepo','Record exact tested runtime/build versions'],
 ['Web interface/backend','Next.js 16.2.0; React 19.2.4','Principal portal and backend workflows'],
 ['Typing and styles','TypeScript 5.9.3; Tailwind CSS 4.2.2','Application implementation'],
 ['Database','PostgreSQL; 16 proposed as isolated test baseline','Verify the deployed LMS version separately'],
 ['Database access','Prisma 7.9.1 and PostgreSQL adapter/driver','Schema access and reviewed migrations'],
 ['Mobile','Expo 55.0.31; React Native 0.83.10; Expo Router','Evaluate selected enabled functions'],
 ['Desktop','Electron 41.0.3; Electron Builder','Native Administrator Console'],
 ['Browsers','Current Chrome/Edge initially; Firefox/Safari where tested','Record exact versions and enabled activities'],
 ['CodeLab execution','Configured Judge0-compatible endpoint','Verify service version, languages, isolation and limits'],
 ['Repository coursework','Git and GitHub integration','Repository authorization is separate from LMS login'],
 ['Files and email','Persistent object storage; configured account email service','Verify effective upload limits and supported delivery'],
 ['Testing','Playwright 1.63.0; TypeScript/Node and harness checks','Fresh functional, access and browser evidence'],
 ['Observation and analysis','Developer tools, logs, database observations; Excel or LibreOffice Calc proposed','Select workload tools and record actual versions']
],[1.3,2.7,2.68])
para('Proposed evaluation targets are Windows 11 x64, Android 10+ and iOS 16+, conditional on verified builds.')
comment(d.paragraphs[-1],'Confirm OS/build targets and dependency versions against the evaluation release. Framework versions here reproduce the inspected development configuration; they are not device test results.')

heading('3.12 Role Based Flowcharts',page=True)
role_texts={
 'student':('Student View','Existing students choose the institute portal and enter email/password; new students use the supported registration process. Backend checks establish account, role and institute access before opening the dashboard. Students access approved coursework and their own grades, personal records, study tools, and authorized collaboration/programming tasks. Calendar access comprises readable saved events and coursework deadlines; faculty-only events and recurring student timetable projection are excluded. Each protected operation is validated by its source workflow, and the session can return to the dashboard or end through logout.'),
 'instructor':('Instructor View','Instructors authenticate through the portal and operate on assigned courses. They manage materials, assessments, submissions, grades and enrollment decisions; inspect descriptive analytics and rule-based attention indicators; and use authorized schedules, calendar and directory functions. Program Head publication is available only with a valid designation. Scheduling changes preserve earlier revisions, while occurrence exceptions affect permitted future classes. Programming results and repository review remain distinct. Moderation, rewards and owned study functions follow their defined permissions.'),
 'administrator':('Administrator View','The Administrator launches the native Console and enters email/password. The backend validates active status and Administrator authority before providing protected operations. Supported functions include accounts, roles, instructor assignments, Super User provisioning, institute/faculty configuration, course and calendar administration, log review, and demonstrated reporting/continuity tools. Audit records are reviewed rather than freely rewritten. Confirmation and reauthentication apply where required by the implemented operation. Additional browser/mobile administration is not assumed by this evaluation.'),
 'super-user':('Super User View','The Super User authenticates through the Web Portal and selects a supported institute, validated independently of the account’s home institute. The academic hub provides permitted faculty information, timetable filters, narrow leadership assignments, shared announcements, and safe course summaries. Supervision excludes student identities, grades, submissions, attempts and answers. The role does not permit grading, assessment participation, course schedule editing, or generic Administrator tools. Owned notes/events remain private. Users return to the hub for another action or log out.')}
for role,(title,text) in role_texts.items():
    heading('3.12.'+str(list(role_texts).index(role)+1)+' '+title,2,page=role!='student')
    figure(role,title+' navigation and access flow',height=6.5 if role=='student' else 7.1);para(text)
heading('3.12.5 Enrollment Request Flow',2,page=True);figure('enrollment','Course enrollment request and instructor decision',height=7.8)
para('A valid join code does not immediately grant course access. The backend checks the institute and existing status, rejects invalid/duplicate requests, and records eligible requests as pending. The student may continue using other modules while waiting. The assigned instructor records approval or rejection; only approval grants enrolled-course access. The dashed review transition represents a later instructor action rather than a synchronous student operation. Expired sessions return to authentication, and failed source operations provide feedback without implying that unsaved changes succeeded.')
heading('3.13 Entity Relationship Diagrams',page=True)
para('Figures 8–11 summarize selected current Prisma relationships and the agreed coordination target. User stores account roles, Enrollment links students and courses, and grading uses applicable submission records and policies. PK/FK denote primary/foreign keys. These scoped figures must agree with the frozen release schema.')
for name,title in [('erd-academic','Selected existing account and academic relationships'),('erd-study','Selected existing study and progression relationships'),('erd-programming-forum','Selected existing programming and forum relationships'),('erd-proposed-coordination','Proposed academic coordination relationships')]:
    figure(name,title,height=6.9 if name=='erd-academic' else 8.35,page=name!='erd-academic')
    if name=='erd-proposed-coordination':
        para('The target model introduces CourseSchedule and CourseScheduleException and extends Institute, User, Course, CalendarEvent and Note. Schedule revisions store weekdays, minutes, date bounds, timezone, room and attribution; exceptions identify original occurrences. Event audiences/time fields and source links support private snapshots. One head per program, head membership, unique conversion links, transactions and version comparisons enforce the agreed integrity rules. Dashed items are proposed until reviewed migrations and feature verification establish them.')
    else:para('Relationships identify the principal references in this functional group. Optional fields permit the stated unassigned or unlinked state; access policy is enforced separately from database relationships. Additional supporting records, indexes, constraints and assessments must be checked against the frozen evaluation schema.')
comment(d.paragraphs[-1],'Confirm these selected ERDs against the frozen release schema. The diagrams separate the current working-tree records from planned coordination additions; migration completion and production schema have not been inferred.')
heading('3.14 Data Flow Diagrams',page=True);figure('dfd-context','Context diagram of authorized CdM LMS information exchange',height=6.4)
para('The context diagram identifies the four role categories as external actors exchanging requests and permitted results with CdM LMS. Repository and execution services support their respective programming workflows. Data exchange occurs through authorized backend operations rather than direct client database access. Inputs, returned records and external-service results are bounded by role, institute, ownership and course relationships.')
figure('dfd-level-one','Level 1 functional data flow and principal record groups',height=8.0,page=True)
para('Level 1 decomposes the system into authentication/authorization, academic coursework and grading, personal study/progression/collaboration, academic coordination, and administration. D1–D5 represent logical record groups rather than additional database tables. D4 includes existing event records and the proposed coordination extension. Generated classes and syllabus deadlines are calendar read outputs, not routine persisted event rows. Each functional process validates its requests and retrieves or changes only authorized records.')
heading('3.15 Release Configuration and Status Record')
para('Before recruitment, the researchers will freeze the release revision and record client builds, schema/migration state, endpoint configuration, enabled languages, repository provider, accepted files, device/browser matrix, network conditions, role accounts, datasets, and known limitations. Implementation and test status will be recorded separately. Secrets will remain outside manuscript examples and public repositories.')
table('Feature status categories for the evaluation record',['Category','Required release evidence'],[
 ['Existing module','Source/workflow identified; actual coverage and role access verified'],
 ['Planned extension','Requirement and target model recorded; not reported as complete'],
 ['Tested function','Fresh expected/actual results, setup and retest status'],
 ['Evaluation enabled','Reviewed function/build included in assigned participant tasks']
],[1.55,5.13])
comment(d.paragraphs[-1] if d.paragraphs[-1].runs else d.paragraphs[-2],'Complete the release record before evaluation. This comparison manuscript includes a proposed research procedure but no invented execution results, participant counts, migration state, or institution-wide deployment claim.')

heading('REFERENCES',0,True)
fullbody='\n'.join(p.text for p in d.paragraphs)
journals=['PLOS ONE','Asian Journal of Education and Social Studies','International Journal of Educational Technology in Higher Education','International Journal of Information Management Data Insights','International Journal in Information Technology in Governance, Education and Business','International Journal of Interactive Mobile Technologies','Frontiers in Education','American Journal of Educational Research','Journal of Computing in Higher Education','International Journal of Computing Sciences Research','Behavioral Sciences','Social Sciences & Humanities Open','Humanities and Social Sciences Letters','Education and Information Technologies','Education Sciences','Decision Sciences Journal of Innovative Education','Psychology and Education: A Multidisciplinary Journal','Journal of Education and Learning','Journal of Applied Structural Equation Modeling','ACM Transactions on Computing Education','Online Learning','Journal of Applied Research in Memory and Cognition','Computer Applications in Engineering Education','GEO Academic Journal','PETIR','Technium Education and Humanities','SN Computer Science']
for ref in REFS:
    if ref['key'] not in fullbody and not (ref['key']=='ISO' and 'ISO' in fullbody):continue
    p=para(ref['text'],False,False,before=0,after=12,spacing=2)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    text=p.text; spans=[]
    for journal in journals:
        if journal in text:
            a=text.index(journal);b=a+len(journal)
            m=re.match(r', (\d+)',text[b:])
            if m:b+=len(m[0])
            spans.append((a,b))
    if spans:
        for r in list(p.runs):r._r.getparent().remove(r._r)
        cursor=0
        for a,b in sorted(spans):
            for part,italic in [(text[cursor:a],False),(text[a:b],True)]:
                rr=p.add_run(part);rr.font.name='Times New Roman';rr.font.size=Pt(12);rr.italic=italic;rr.font.color.rgb=RGBColor(0,0,0)
            cursor=b
        rr=p.add_run(text[cursor:]);rr.font.name='Times New Roman';rr.font.size=Pt(12);rr.font.color.rgb=RGBColor(0,0,0)
    p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
    p.paragraph_format.keep_together=True

settings=d.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
d.save(ROOT/'generated.docx')

# Preserve opaque/source-owned OOXML parts and original media byte for byte.
with ZipFile(SRC) as zin,ZipFile(ROOT/'generated.docx') as gen:
    original={n:zin.read(n) for n in zin.namelist()};generated={n:gen.read(n) for n in gen.namelist()}
    editable={'word/document.xml','word/_rels/document.xml.rels','word/header1.xml','word/_rels/header1.xml.rels','word/settings.xml','[Content_Types].xml'}
    final=dict(original)
    for name,data in generated.items():
        if name in editable or name not in original:final[name]=data
    with ZipFile(FINAL,'w',ZIP_DEFLATED) as z:
        for name,data in final.items():z.writestr(name,data)
    preserved=[n for n in original if n not in editable]
    assert all(final[n]==original[n] for n in preserved)
manifest={'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'output':str(FINAL),'preserved_parts':preserved,'figures':fig_num,'tables':table_num,'paragraphs':len(d.paragraphs),'words':len(fullbody.split()),'reference_count':sum(ref['key'] in fullbody for ref in REFS)}
(ROOT/'build-manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
(ROOT/'revised-text.txt').write_text('\n'.join(p.text for p in d.paragraphs),encoding='utf-8')
print(json.dumps(manifest,ensure_ascii=False))
