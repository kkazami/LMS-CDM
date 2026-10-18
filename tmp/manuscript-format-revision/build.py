from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import json, re, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parent
SRC=Path(r'C:\Users\kirby\Downloads\LMS FINAL (1).docx')
OUT=Path(r'C:\Users\kirby\Desktop\lms-monorepo-skeleton\output\documents')
OUT.mkdir(parents=True,exist_ok=True)
FINAL=OUT/'CdM_LMS_Chapters_1_to_3_Format_Revised.docx'
source=Document(SRC);d=Document(SRC)
C=json.loads((ROOT/'content.json').read_text(encoding='utf-8'))
REFS=json.loads((ROOT/'references.json').read_text(encoding='utf-8'))
source_paras=[p.text.strip() for p in source.paragraphs]
cover=[deepcopy(p._p) for p in source.paragraphs[:26]]
for child in list(d._element.body):
    if child.tag!=qn('w:sectPr'):d._element.body.remove(child)
for el in cover:d._element.body.insert(len(d._element.body)-1,el)
cover_title=d.paragraphs[3]
for r in list(cover_title.runs):r._r.getparent().remove(r._r)
r=cover_title.add_run(C['title']);r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=True
cover_title.style=d.styles['Title'];cover_title.alignment=WD_ALIGN_PARAGRAPH.CENTER

datep=next(p for p in d.paragraphs if p.text.strip()=='May 2026')
d.add_comment(runs=datep.runs,text='The supplied cover date has been preserved. Confirm the actual submission date before filing. Later-2026 sources whose dates would conflict with this cover were excluded from the revised literature.',author='Manuscript Review',initials='MR')

def clean(text):
    text=text.replace('technical checks by qualified IT experts','technical checks conducted by the researchers').replace('BSCPE','BSCpE').replace('Pan et al. (2022)','Pan et al. (2023)').replace('Jack et al. (2025)','Jack (2025)')
    verbs={'describes':'describe','emphasizes':'emphasize','identifies':'identify','provides':'provide','supports':'support'}
    text=re.sub(r'(et al\. \(\d{4}\)) (describes|emphasizes|identifies|provides|supports)\b',lambda m:m[1]+' '+verbs[m[2]],text)
    text=text.replace('points and leaderboards in a flipped statistics classroom and discussed benefits','points and leaderboards in a flipped statistics classroom and considered benefits')
    text=text.replace('Peralta (2024), Guico (2026) further','Peralta (2024) and Guico (2026) further')
    return text
def para(text='',bold=False,indent=True,style='Normal',before=12,after=12,spacing=2.0):
    p=d.add_paragraph(style=style);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    pf=p.paragraph_format;pf.first_line_indent=Inches(.5) if indent else Inches(0)
    pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=spacing;pf.widow_control=True
    r=p.add_run(clean(text));r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0);r.bold=bold
    r._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Times New Roman')
    return p
def heading(text,level=1,page=False):
    p=para(text,True,False,'Title' if level==0 else f'Heading {min(level,3)}',before=18,after=10,spacing=1.0)
    p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True;p.paragraph_format.page_break_before=page
    if level==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p
def comment(p,text):d.add_comment(runs=p.runs,text=text,author='Manuscript Review',initials='MR')
def section(start,end):
    i=next(i for i,t in enumerate(source_paras) if t.startswith(start));j=next((j for j in range(i+1,len(source_paras)) if source_paras[j].startswith(end)),len(source_paras))
    return [t for t in source_paras[i+1:j] if t]
def paragraphs(items):
    return [para(t) for t in items]
fig_num=0
def figure(name,title,width=6.4,height=7.6,page=False):
    global fig_num
    fig_num+=1
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pf=p.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(8);pf.line_spacing=1;pf.keep_with_next=True;pf.page_break_before=page
    path=ROOT/'figures'/f'{name}.png';im=Image.open(path);ratio=im.height/im.width
    w=min(width,height/ratio)
    pic=p.add_run().add_picture(str(path),width=Inches(w));pic._inline.docPr.set('descr',title)
    cap=para(f'Figure {fig_num}',False,False,before=0,after=5,spacing=1)
    cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.keep_with_next=True
    text=para(title,False,False,before=0,after=12,spacing=1.3)
    text.alignment=WD_ALIGN_PARAGRAPH.CENTER;text.paragraph_format.keep_together=True
    for q in [cap,text]:
        for r in q.runs:r.italic=True;r.font.size=Pt(11)
    return cap
table_num=0
def table(title,headers,rows,widths):
    global table_num
    table_num+=1
    cap=para(f'Table {table_num}. {title}',False,False,before=10,after=8,spacing=1.0)
    cap.paragraph_format.keep_with_next=True
    for r in cap.runs:r.italic=True;r.font.size=Pt(11)
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for vals in [headers]+rows:
        cells=t.rows[0].cells if vals is headers else t.add_row().cells
        for i,(cell,text) in enumerate(zip(cells,vals)):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');marg.append(e)
            tcpr.append(marg)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'B8B8B8');borders.append(e)
            tcpr.append(borders)
            if vals is headers:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');tcpr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_with_next=False
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(clean(str(text)));r.font.name='Times New Roman';r.font.size=Pt(10.5);r.bold=vals is headers;r.font.color.rgb=RGBColor(0,0,0)
        trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if vals is headers:trpr.append(OxmlElement('w:tblHeader'))
    # Keep a table header and its first row together, without keeping the entire table.
    for p in t.rows[0].cells[0].paragraphs:p.paragraph_format.keep_with_next=True
    return cap

# Preserve the exact branded header appearance shown in LMS FINAL.pdf.
sec=d.sections[0];sec.top_margin=Inches(1.70);sec.header_distance=Inches(20/72)
header=sec.header
for el in list(header._element):header._element.remove(el)
hp=header.add_paragraph();hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
hp.paragraph_format.space_before=Pt(0);hp.paragraph_format.space_after=Pt(0);hp.paragraph_format.line_spacing=1
hp.add_run().add_picture(str(ROOT/'source-header.png'),width=Inches(6.6875),height=Inches(1.365))

heading('CHAPTER 1: INTRODUCTION',0,True)
heading('1.1 BACKGROUND OF THE STUDY')
background=C['background']
paragraphs(background)
heading('1.2 STATEMENT OF THE PROBLEM')
for text in section('1.2 STATEMENT','1.3 Statement'):
    text=text.replace('What is the level of the LMS in supporting','What is the evaluated software quality of CdM LMS in supporting').replace('What is the level of the system in','What is the evaluated software quality of CdM LMS in')
    para(text,indent=not bool(re.match(r'^\d\.|^•',text)))
heading('1.3 STATEMENT OF THE OBJECTIVES')
heading('1.3.1 General Objective',2);para(section('1.3.1 General','1.3.2 Specific')[0])
heading('1.3.2 Specific Objectives',2)
for text in section('1.3.2 Specific','1.4 CONCEPTUAL'):
    if text.startswith('The study'):continue
    if re.match(r'^\d+\.',text):
        text=text.replace('verified and approved accounts','account-status controls').replace('Develop the proposed faculty directory','Develop the faculty directory')
        para(text,indent=False)
heading('1.4 CONCEPTUAL FRAMEWORK',page=True);figure('conceptual','Conceptual framework of CdM LMS using IPO with feedback',height=8.0)
paragraphs(C['conceptual'])
heading('1.5 OPERATIONAL FRAMEWORK');figure('agile','Adapted Agile development cycle for CdM LMS',height=4.2)
paragraphs(C['operational'])
heading('1.6 SCOPE AND LIMITATIONS');paragraphs(C['scope'])
heading('1.7 SIGNIFICANCE OF THE STUDY')
for text in section('1.7 SIGNIFICANCE','1.8 GENERAL'):
    text=text.replace('The developed Colegio de Montalban Learning Management System (CdM LMS) is significant','The proposed CdM LMS may contribute').replace('as it provides','by providing')
    if text.startswith('Academic supervisors'):
        text='Super Users and Program Heads. Restricted Super User supervision may support authorized academic oversight. Program Heads may benefit from faculty information and designation-specific event publication. These responsibilities will remain distinct from teaching and system administration.'
    para(text)
heading('1.8 GENERAL FEATURES OF THE SYSTEM / STUDY')
p=para('CdM LMS comprises complementary applications connected through shared backend services and a centralized database. The following groups describe intended feature coverage. A release record will distinguish existing modules, planned extensions, tested functions, and functions enabled for evaluation. Availability is governed by role and platform coverage.')
comment(p,'Before evaluation, replace the proposed release-status record with confirmed module/build evidence. The academic coordination extension remains planned in the inspected repository; its inclusion here is a development requirement, not a completed result.')
for category,groups in C['features_by_category']:
    heading(category,2)
    for title,text in groups:
        heading(title,3);para(text)
table('Role and platform responsibilities',['Role','Principal interface','Permitted responsibility','Boundary'],[
 ['Student','Web; selected mobile functions','Approved coursework, own grades, study tools and collaboration','No faculty directory or supervision'],
 ['Instructor','Web; selected mobile functions','Assigned teaching, grading, enrollment and schedules','Program Head capabilities require designation'],
 ['Super User','Web academic operations hub','Safe supervision and narrow faculty coordination','No grading, assessments, schedule editing or system administration'],
 ['Administrator','Native Desktop Console','Protected account/configuration, audit and supported continuity operations','Additional browser/mobile administration is outside this evaluation scope']
],[1.0,1.55,2.5,1.63])
heading('1.9 DEFINITION OF TERMS')
definitions=section('1.9 DEFINITION','CHAPTER 2')
definitions=[t for t in definitions if not t.startswith('The following')]
definitions=[t.replace('Spaced-Repetition-Oriented Review - Flashcard review organized to support repeated study over time, without assuming an unverified scheduling algorithm.','Flashcard Review - Repeated study of saved question-and-answer cards using the review and mastery functions available in the evaluated release.') for t in definitions]
definitions.extend(['BSIT - Bachelor of Science in Information Technology, one of the ICS programs within the study context.','BSCpE - Bachelor of Science in Computer Engineering, one of the ICS programs within the study context.','Dean - An academic leadership designation represented in the institute directory; the designation itself does not grant additional system permissions.'])
for text in sorted(definitions,key=lambda t:t.casefold()):para(text,indent=False,before=6,after=6)

heading('CHAPTER 2: REVIEW OF RELATED LITERATURE',0,True)
intro=section('CHAPTER 2','2.1.1 Multi')
for text in intro:
    text=text.replace('seven topics','eight topic groups')
    if text.startswith('The project developers reviewed'):
        text=text.replace('and Online Collaboration and Knowledge Sharing.','Online Collaboration and Knowledge Sharing; and Academic Coordination and Information Boundaries.')
    para(text)
heading('2.1 Related Literature')
start=source_paras.index(next(t for t in source_paras if t.startswith('2.1.1 Multi')))
end=source_paras.index(next(t for t in source_paras if t.startswith('CHAPTER 3')))
inserted=False
for text in source_paras[start:end]:
    if not text:continue
    if text.startswith('2.2 ') and not inserted:
        heading('2.1.8 Academic Coordination and Information Boundaries',2)
        para('Institutional LMS support extends beyond individual coursework. Guico (2026) identifies faculty preparation, resource access, and continuing training as management concerns. Llamas et al. (2025) identifies monitoring, assessment, communication, and course organization as user-experience concerns. Together, these findings provide reasons to examine faculty coordination and access to appropriate course information. They do not validate the particular schedule schema or supervision policy proposed for CdM LMS.')
        para('The project’s academic coordination requirements will therefore be established through local consultation and explicit role policies. Directory information, timetable revisions, event audiences, and restricted supervision will be tested against those requirements. Academic designation and technical privilege will remain separate; shared calendar visibility will not imply access to private notes or individual assessment records.')
        inserted=True
    if text.startswith('2.5 Synthesis'):text=text.replace('2.5','2.4',1)
    if text.startswith('2.6 Research'):text=text.replace('2.6','2.5',1)
    if re.match(r'^2\.\d',text):heading(text,2 if text.count('.')>=2 else 1);continue
    if 'Bernardez et al.' in text:continue
    if 'Francisco (2026)' in text:
        if text.startswith('Institutional support'):
            text=text.split('Francisco (2026)')[0].strip()+' These findings place course management within faculty support and institutional coordination.'
        elif text.startswith('The reviewed Philippine'):text='The reviewed Philippine studies inform the project’s computing-course, LMS, programming, and collaborative workflows. Their findings do not establish conditions within ICS. Local requirements gathering will identify the needs and experiences of the intended users.'
        else:text=text.replace(', and Francisco (2026)','').replace(' and Francisco (2026)','')
    text=text.replace('these five studies','the reviewed studies').replace('These five studies','The reviewed studies').replace('The five studies','The reviewed studies').replace('Across these six studies','Across the reviewed studies').replace('Together, the five studies','Together, the reviewed studies').replace('The standard is a supporting framework rather than one of the five studies counted within each topic.','The standard supports the quality criteria used in this study.')
    if 'Jack et al. (2025)' in text:text=text.replace('Jack et al. (2025) examined','Jack (2025), in a conference abstract, discussed')
    if text.startswith('The reviewed studies address digital learning'):
        chunks=text.split('A second part');para(chunks[0].strip())
        if len(chunks)>1:
            rest='A second part'+chunks[1];chunks=rest.split('A third part');para(chunks[0].strip())
            if len(chunks)>1:para('A third part'+chunks[1])
    else:para(text)

heading('CHAPTER 3: METHODOLOGY',0,True)
heading('3.1 RESEARCH DESIGN');paragraphs(C['design'])
heading('3.2 SYSTEM AND PROJECT OVERVIEW');paragraphs(C['overview'])
heading('3.3 RESPONDENTS');ps=paragraphs(C['respondents'])
comment(ps[-1],'The group reports a historical prior-semester ICS student population of 1,602: 932 IT students and 670 CpE students. Update these figures from an authorized current enrollment record and state its semester and date. These historical population figures are not the target sample or completed respondent count. Instructor population, full-time/part-time breakdown, and Student/Instructor target samples are not yet confirmed; supply them after adviser approval.')
heading('3.4 SAMPLING TECHNIQUE');paragraphs(C['sampling'])
heading('3.5 INSTRUMENTS AND METHODS FOR DATA COLLECTION')
heading('3.5.1 Questionnaires and Instrument Review',2);paragraphs(C['instruments'])
table('Role appropriate respondent tasks',['Respondent','Platform','Representative tasks'],[
 ['ICS Students','Web and selected mobile functions','Approved coursework, own grades, notes/tasks, study tools and Knowledge Exchange'],
 ['ICS Instructors','Web and selected mobile functions','Assigned materials, submissions, grading, enrollment decisions and schedules'],
 ['Instructors designated as Program Heads','Web; permitted designation functions','Authorized faculty or program event publication; part of the Instructor group']
],[1.55,1.70,3.43])
heading('3.5.2 Local Requirements Gathering',2);ps=paragraphs(C['requirements'])
comment(ps[0],'Confirm the consultation and recruitment arrangements with the adviser before proceeding. These paragraphs describe planned methods, not completed interviews or approved institutional access.')
heading('3.5.3 Technical Testing and Observation',2);ps=paragraphs(C['testing'])
comment(ps[1],'Approve dataset size, active sessions, request rate, simultaneous CodeLab executions, repetitions, timing thresholds and measurement tools before testing. The researchers perform these checks; test accounts do not create additional respondent groups.')
table('Evidence for selected quality characteristics',['Characteristic','Researcher checks','Respondent evidence'],[
 ['Functional Suitability','Expected/actual computations, submissions, schedules and conversions','Ratings of directly used functions'],
 ['Interaction Capability','Task completion, navigation and feedback observations','Student/Instructor task ratings'],
 ['Security','Prohibited requests and private-record isolation','Observable permissions and own-record access'],
 ['Performance Efficiency','Recorded workloads, timings and resource observations','Perceived responsiveness on performed tasks'],
 ['Reliability','Interrupt/retry, duplicate or stale saves, persistence','Encountered failures during the session'],
 ['Compatibility','Permitted shared records and native-console co-existence','Client tasks within the respondent’s access']
],[1.5,2.65,2.53])
heading('3.5.4 Data Analysis and Statistical Treatment',2)
ps=[]
for i,text in enumerate(C['analysis']):
    if i!=1:
        ps.append(para(text));continue
    intro=text.split('Equivalently,')[0].strip()
    ps.append(para(intro))
    # Native editable Word fraction and frequency subscripts.
    eqp=d.add_paragraph();eqp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    eqp.paragraph_format.space_before=Pt(6);eqp.paragraph_format.space_after=Pt(6)
    om=OxmlElement('m:oMathPara');o=OxmlElement('m:oMath');om.append(o)
    def mr(container,text):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);container.append(r)
    def frequency(container,num):
        sub=OxmlElement('m:sSub');base=OxmlElement('m:e');mr(base,'f');lower=OxmlElement('m:sub');mr(lower,str(num));sub.extend([base,lower]);container.append(sub)
    mr(o,'WM = ');frac=OxmlElement('m:f');num=OxmlElement('m:num');den=OxmlElement('m:den')
    for j,n in enumerate([4,3,2,1]):
        if j:mr(num,' + ');mr(den,' + ')
        if n>1:mr(num,str(n))
        frequency(num,n);frequency(den,n)
    frac.extend([num,den]);o.append(frac);eqp._p.append(om)
    para('Here, WM is the frequency-weighted item mean and the subscripted frequencies count valid responses at scores 4, 3, 2 and 1. Frequency weighting represents response counts rather than different importance weights for users or quality characteristics. Response distributions and valid counts will accompany means; median summaries may also be reported.')
comment(ps[1],'Confirm the scoring and aggregation procedure with the adviser/statistician. Report valid counts and response distributions. Do not invent interpretation bands or treat N/A as a numeric rating.')
heading('3.5.5 Participant Protection and Information Handling',2);ps=paragraphs(C['ethics'])
comment(ps[-1],'Finalize institutional permission, retention period, storage location, access list and withdrawal/deletion procedures. These planned measures do not assert that approval has already been granted.')
heading('3.5.6 Evaluation Release Record',2)
para('Before recruitment, the researchers will record the application revision, client builds, schema and migration state, enabled modules and languages, repository provider, file limits, device/browser matrix, network conditions, role accounts, datasets, and known limitations. Existing modules, planned extensions, tested functions, and functions enabled for respondent tasks will be identified separately. Administrator and Super User functions will be verified by the researchers with controlled accounts; ordinary respondents will not rate functions outside their authority. Secrets and identifiable academic records will be excluded from manuscript examples.')
heading('3.6 HARDWARE REQUIREMENTS')
para('Equipment specifications below are planning estimates for development and a small evaluation pilot. They are not experimentally established minimums or verified concurrent capacity. The actual device and hosting configuration will be recorded before task sessions.')
table('Estimated development and client hardware',['Equipment','Estimated baseline','Recommended','Purpose'],[
 ['Development workstation','4 cores; 16 GB RAM; 50 GB free SSD','6+ cores; 32 GB RAM; 100 GB free SSD','Build clients and run local checks'],
 ['Web computer','Dual core; 4 GB RAM; 1366×768; 1 GB free','4 cores; 8 GB RAM; 1080p; 5 GB free','Academic and coordination tasks'],
 ['Android device','Build-compatible; about 4 GB RAM; 1 GB free','6+ GB RAM; 2 GB free','Enabled mobile tasks'],
 ['iOS device if evaluated','Compatible iPhone/iPad; 1 GB free','Supported device; 2 GB free','Verified iOS build only'],
 ['Administrator workstation','64-bit dual core; 4 GB RAM; 2 GB free','4 cores; 8 GB RAM; 5 GB free','Native Console and exports'],
 ['Network equipment','Reliable router/AP and connectivity','Wired laboratory links and sufficient Wi-Fi capacity','Connect authorized clients/services']
],[1.45,1.8,1.8,1.63])
para('Storage allowances cover dependencies, temporary files, cache, and downloads rather than measured installer sizes. Emulators, containers, and datasets may need more development space. Enabled 3D activities require a tested WebGL 2-capable browser/graphics configuration; dedicated graphics is not assumed for ordinary LMS tasks.')
table('Proposed pilot hosting and network configuration',['Component','Initial estimate','Condition to verify'],[
 ['Application service','2 vCPU / 4 GB RAM on a conventional server','For serverless hosting, record plan/function limits instead'],
 ['PostgreSQL','2 vCPU / 4 GB RAM; 10–20 GB initial records/index storage','Connection limits, growth and query behavior'],
 ['Persistent file storage','20–50 GB initial object-storage allowance','Materials, submissions, snapshots and retention'],
 ['Code execution if self-hosted','Separate 4 vCPU / 8 GB RAM; 30–50 GB storage','Languages, isolation, test limits and execution concurrency'],
 ['Backups','Separate protected storage','Retention capacity and demonstrated restoration'],
 ['Individual connection','About 5 Mbps down / 2 Mbps up; prefer 10/5','Measured latency, contention and file sizes'],
 ['Shared lab connection','About 100 Mbps down / 20–50 Mbps up for an initial 20–30-device session','Actual simultaneous workload and connection stability']
],[1.45,2.55,2.68])
para('The initial sizing assumption is approximately 25–50 simultaneously active ordinary-task users, not a research sample size or proven capacity. Code executions will be limited and measured separately. Lower bandwidth may support basic retrieval, while uploads and concurrent activity may need more. Managed-service quotas and the demonstrated workload take precedence over these estimates.')
heading('3.7 SOFTWARE REQUIREMENTS')
para('The software record identifies the current development stack and proposed evaluation targets. It will be frozen against the actual release revision; framework support alone does not establish CdM client compatibility.')
table('Software environment and release configuration',['Software/service','Recorded choice','Purpose/qualification'],[
 ['Development OS','Supported Windows/macOS/Linux environment','Platform-specific build requirements apply'],
 ['Node and workspace tools','Supported Node 22 LTS patch ≥22.12; pnpm 10.32.1; Turborepo','Record exact tested runtime/build versions'],
 ['Web interface/backend','Next.js 16.2.0; React 19.2.4','Principal portal and backend workflows'],
 ['Typing and styles','TypeScript 5.9.3; Tailwind CSS 4.2.2','Application implementation'],
 ['Database','PostgreSQL; 16 proposed as isolated test baseline','Verify the deployed LMS version separately'],
 ['Database access','Prisma 7.9.1 and PostgreSQL adapter/driver','Schema access and reviewed migrations'],
 ['Mobile','Expo 55.0.31; React Native 0.83.10; Expo Router','Evaluate selected enabled functions'],
 ['Desktop','Electron 41.0.3; Electron Builder','Native Administrator Console'],
 ['Browsers','Current Chrome/Edge initially; Firefox/Safari where tested','Record exact versions and enabled activities'],
 ['CodeLab execution','Configured Judge0-compatible endpoint','Verify service version, languages, isolation and limits'],
 ['Repository coursework','Git and GitHub integration','Repository authorization is separate from LMS login'],
 ['Files and email','Persistent object storage; configured account email service','Verify effective upload limits and supported delivery'],
 ['Testing','Playwright 1.63.0; TypeScript/Node and harness checks','Fresh functional, access and browser evidence'],
 ['Observation and analysis','Developer tools, logs, database observations; Excel or LibreOffice Calc proposed','Select workload tools and record actual versions']
],[1.3,2.7,2.68])
para('Proposed evaluation targets are Windows 11 x64, Android 10+ and iOS 16+, conditional on verified builds.')
comment(d.paragraphs[-1],'Confirm OS/build targets and dependency versions against the evaluation release. Framework versions here reproduce the inspected development configuration; they are not device test results.')


heading('3.8 FLOWCHART',page=True)
para('The diagrams separate access and registration, role navigation, and course enrollment details. Oval terminators identify the start or end; rectangles identify processes; parallelograms identify input/output; diamonds identify decisions; and rectangles with double side lines identify detailed subprocesses. Circles continue a flow on the same page, while pentagonal connectors continue it on another diagram. Matching letters identify the destination. Each protected operation remains subject to backend authorization.')
figure('symbol-key','Symbols used in the system flowcharts',height=3.5)
para('Off-page connectors R and L link registration with the access chart; S, I, A and U link to Student, Instructor, Administrator and Super User overviews. E and E2 lead to the enrollment request and Instructor review details. Same-page pairs L1, R1 and E1 allow correction and retry; S1, I1, A1 and U1 return to the corresponding dashboard. A dashed line marks the later review of a recorded pending request, which does not block the Student from using other modules.')
heading('3.8.1 Account Access and Registration',2,page=True)
figure('access','Account access and role routing',height=6.55)
para('Users open the appropriate client and authenticate using email, password and institute selection. The backend validates credentials, account status, role and permitted institute before establishing a session. Invalid access produces feedback and a retry through L1. The target release routes successful access to the appropriate role overview. New public Student or Instructor accounts use R; privileged accounts are created through authorized administrative operations. The Super User routing branch represents the agreed expansion and must be verified in the evaluation release.')
figure('registration','Public Student and Instructor registration detail',height=7.1,page=True)
para('The registration process validates required fields, institute, permitted public role and duplicate email. Invalid input returns through R1 for correction. Valid registration hashes the password, creates the account, and displays confirmation. L returns to the access chart. Any email verification or account-approval steps must agree with the evaluated registration workflow.')
heading('3.8.2 Student View',2,page=True);figure('student','Student overview with links to detailed enrollment',height=7.15)
para('S opens the Student Dashboard after successful access or a completed enrollment action. Students select approved coursework and assigned deadlines, request enrollment, view their own grades, and use permitted study, calendar, coding and collaboration tools. Personal Tasks remain separate from the academic To-Do list. Each ordinary module returns through S1; Join Courses continues through E. Students do not gain faculty-directory or supervisory privileges. Log out ends the session.')
figure('enrollment','Student course enrollment request detail',height=7.2,page=True)
para('E receives the Join Courses action. The backend checks the code, institute and existing enrollment status. Invalid or duplicate requests produce feedback and retry through E1. An eligible request is saved as PENDING, then shown to the Student. S returns to the Student overview without granting course access. E2 identifies the Instructor’s later review of the saved request.')
figure('enrollment-review','Instructor enrollment decision detail',height=7.1,page=True)
para('The assigned Instructor opens the request queue and the backend checks authority and the current request state. Approval saves APPROVED and grants enrolled-course access; rejection saves REJECTED and does not grant access. Updated status becomes available through the Student’s permitted course view. I returns to the Instructor overview. A valid course code alone does not enroll the Student immediately.')
heading('3.8.3 Instructor View',2,page=True);figure('instructor','Instructor overview with enrollment review continuation',height=7.15)
para('I opens the Instructor Dashboard. Instructors manage assigned courses and resources, configure grades, inspect descriptive analytics, review enrollment requests through E2, and assess CodeLab results or Git-linked work through their separate workflows. Calendar and faculty-directory actions follow their permissions; Program Head publishing requires that designation. Moderation, verification and owned study functions remain bounded by their rules. I1 returns to the dashboard, while Log out ends the session.')
heading('3.8.4 Administrator View',2,page=True);figure('administrator','Native Administrator Console overview',height=7.15)
para('A opens the native Administrator Console after validated Administrator access. Supported actions include accounts and roles, academic configuration, schedule/event administration, audit review, reports and continuity functions. Only authorized Administrators provision or revoke Super User accounts. Historical audit records are reviewed rather than freely rewritten. A1 returns to the dashboard. The researchers will verify these restricted functions with controlled accounts; no separate Administrator respondent group is introduced.')
heading('3.8.5 Super User View',2,page=True);figure('super-user','Restricted Super User academic operations overview',height=7.15)
para('U opens the web-based academic operations hub for an authorized institute. The Super User uses permitted faculty and leadership information, reads timetables, publishes permitted shared events, inspects safe course summaries, and manages owned private records. Supervision excludes student identities, individual grades, submissions, attempts and answers. The role does not permit grading, assessment participation, schedule editing or generic administration. U1 returns to the hub; Log out ends the session. This expansion will undergo researcher-conducted technical checks.')
heading('3.9 ENTITY RELATIONSHIP DIAGRAM',page=True)
para('The ERDs show selected current Prisma relationships and separate proposed coordination additions. User stores account roles. Repeated entities denote the same table. PK, FK and UK identify primary, foreign and unique keys. Crow’s-foot notation uses a circle for zero, a bar for one and a fork for many; optional child collections may contain zero records.')
for name,title,text in [
 ('erd-academic','Account, institute and course enrollment relationships','Institute links User and Course records. Course.instructorId is optional and references a User acting as Instructor. Enrollment links a Student User to a Course and stores pending, approved or rejected status under a unique course/student pair. Session records belong to User. Role and resource authorization are enforced separately from these foreign-key links.'),
 ('erd-grading','Coursework, submission and grading relationships','Course links SyllabusItem records and an optional GradingPolicy. CategoryWeight records define the policy’s component weights. StudentSubmission connects an item to a Student User, with a unique item/student pair and recorded grade or status. Computed grade summaries use the applicable grading rules; they are not modelled as a single undifferentiated Grade table.'),
 ('erd-support','Selected personal ownership and learning activity relationships','User owns Note and CalendarEvent records, creates FlashcardDeck records, and creates ActivityTemplate records through createdBy. CalendarEvent may optionally reference Course, while ActivityTemplate requires a courseId. These selected relationships omit supporting task, card, progress, study-session and gamification tables for readability; their fields and constraints remain part of the release schema.'),
 ('erd-coursework-forum','Separate programming, repository and forum record relationships','ActivityTemplate links ActivitySubmission records containing scores and execution-related outcomes. GithubRepoLink links GithubSubmission records, which also reference the submitting User and SyllabusItem. KxPost links KxAnswer records with an author and optional verifying faculty User. These groups preserve the distinction between code execution, Git coursework and moderated discussion.'),
 ('erd-proposed-coordination','Proposed coordination schema and private snapshot relationships','Dashed outlines distinguish the target additions from current schema records. CourseSchedule and CourseScheduleException will preserve effective-dated revisions and occurrence exceptions. Course and User will receive program and coordination fields. CalendarEvent will receive time, audience, source-note and version fields; Note will receive an event-source link with a unique creator/event combination. Institute.deanId and other membership/attribution links will be checked against the final migration. Snapshot conversions are one-time copies, not ongoing synchronization.')
]:
    figure(name,title,height=5.75 if name=='erd-academic' else 7.05,page=name!='erd-academic');para(text)
comment(d.paragraphs[-1],'Verify the selected ERDs and proposed fields against the final approved migration and evaluation schema. Current code does not establish completion of the coordination extension. This is a scoped relational view, not an exhaustive table dictionary.')
heading('3.10 DATA FLOW DIAGRAM',page=True)
figure('dfd-context','Context diagram of the LMS and external actors',height=6.3)
para('The context diagram presents CdM LMS as one process exchanging data with Students, Instructors, Super Users, Administrators, the repository provider and the code-execution service. Rectangles represent external entities; the rounded process represents the system; and labelled arrows identify the direction of requests or returned information. Own records, safe summaries and administrative results remain limited by their respective permissions. Repository data and execution results are distinct exchanges. The role actors describe system users, not additional research respondent categories.')
heading('3.10.1 Level 1 Functional Data Flow',2,page=True)
figure('dfd-level-one','Level 1 functional data flow and logical record stores',height=7.15)
para('Level 1 separates four processes. Repeated actors denote the same users; D1–D4 are logical record stores. Each process checks applicable permissions. D4 includes proposed coordination records. Recurring classes and deadlines are generated calendar outputs, with Student timetable projection excluded. The final flows will match the evaluation release.')

heading('REFERENCES',0,True)
fullbody='\n'.join(p.text for p in d.paragraphs)
journals=['PLOS ONE','Asian Journal of Education and Social Studies','International Journal of Educational Technology in Higher Education','International Journal of Information Management Data Insights','International Journal in Information Technology in Governance, Education and Business','International Journal of Interactive Mobile Technologies','Frontiers in Education','American Journal of Educational Research','Journal of Computing in Higher Education','International Journal of Computing Sciences Research','Behavioral Sciences','Social Sciences & Humanities Open','Humanities and Social Sciences Letters','Education and Information Technologies','Education Sciences','Decision Sciences Journal of Innovative Education','Psychology and Education: A Multidisciplinary Journal','Journal of Education and Learning','Journal of Applied Structural Equation Modeling','ACM Transactions on Computing Education','Online Learning','Journal of Applied Research in Memory and Cognition','Computer Applications in Engineering Education','GEO Academic Journal','PETIR','Technium Education and Humanities','SN Computer Science']
for ref in REFS:
    if ref['key'] not in fullbody and not (ref['key']=='ISO' and 'ISO' in fullbody):continue
    p=para(ref['text'],False,False,before=0,after=12,spacing=2)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    text=p.text; spans=[]
    for journal in journals:
        if journal in text:
            a=text.index(journal);b=a+len(journal)
            m=re.match(r', (\d+)',text[b:])
            if m:b+=len(m[0])
            spans.append((a,b))
    if not spans:
        for title in ['An empirical study evaluating the effects of gamification in an e-learning environment','Systems and software engineering—Systems and software Quality Requirements and Evaluation (SQuaRE)—Product quality model','Descriptive research: Definition, types, methods and examples','8 Google Classroom tips every teacher should know','Moodle LMS 4.5: Unlock the power of AI for your courses','Proceedings of the 30th International Conference on Computers in Education','Proceedings of the 2024 ACM Virtual Global Computing Education Conference','Proceedings of the 46th International Conference on Software Engineering: Software Engineering Education and Training']:
            if title in text:spans.append((text.index(title),text.index(title)+len(title)))
    if spans:
        for r in list(p.runs):r._r.getparent().remove(r._r)
        cursor=0
        for a,b in sorted(spans):
            for part,italic in [(text[cursor:a],False),(text[a:b],True)]:
                rr=p.add_run(part);rr.font.name='Times New Roman';rr.font.size=Pt(12);rr.italic=italic;rr.font.color.rgb=RGBColor(0,0,0)
            cursor=b
        rr=p.add_run(text[cursor:]);rr.font.name='Times New Roman';rr.font.size=Pt(12);rr.font.color.rgb=RGBColor(0,0,0)
    p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
    p.paragraph_format.keep_together=True

settings=d.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
d.save(ROOT/'generated.docx')

# Preserve opaque/source-owned OOXML parts and original media byte for byte.
with ZipFile(SRC) as zin,ZipFile(ROOT/'generated.docx') as gen:
    original={n:zin.read(n) for n in zin.namelist()};generated={n:gen.read(n) for n in gen.namelist()}
    editable={'word/document.xml','word/_rels/document.xml.rels','word/header1.xml','word/_rels/header1.xml.rels','word/settings.xml','[Content_Types].xml'}
    final=dict(original)
    for name,data in generated.items():
        if name in editable or name not in original:final[name]=data
    with ZipFile(FINAL,'w',ZIP_DEFLATED) as z:
        for name,data in final.items():z.writestr(name,data)
    preserved=[n for n in original if n not in editable]
    assert all(final[n]==original[n] for n in preserved)
manifest={'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'output':str(FINAL),'preserved_parts':preserved,'figures':fig_num,'tables':table_num,'paragraphs':len(d.paragraphs),'words':len(fullbody.split()),'reference_count':sum(ref['key'] in fullbody for ref in REFS)}
(ROOT/'build-manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
(ROOT/'revised-text.txt').write_text('\n'.join(p.text for p in d.paragraphs),encoding='utf-8')
import shutil
shutil.copy2(ROOT/'CdM_LMS_Editable_Diagrams.drawio',OUT/'CdM_LMS_Editable_Diagrams.drawio')
print(json.dumps(manifest,ensure_ascii=False))
